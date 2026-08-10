from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.excel_agent.excel_tools import ExcelAgentTools
from app.excel_agent.tool_executor import ExcelToolExecutor
from app.excel_agent.tool_registry import create_excel_tool_registry
from app.llm.domain import ToolCall
from app.ras_audit.account_mapping import load_ras_ledger_account_mappings
from app.ras_audit.audit_report import RasAuditReportGenerator
from app.ras_audit.fact_context import RasFactContextAttestor, RasFactExtraction
from app.ras_audit.persisted_report import PersistedRasAuditReportService
from app.ras_audit.persistence import SqlAlchemyRasAuditRepository
from app.ras_audit.report_exports import (
    render_ras_report_pdf,
    render_ras_report_xlsx,
)

EXPECTED_SHEETS = (
    "Synthese",
    "Anomalies",
    "Dossiers incomplets",
)


def main() -> None:
    container_root = Path("/workspace")
    repository_root = (
        container_root
        if (container_root / "docs").is_dir()
        else Path(__file__).resolve().parents[2]
    )
    docs_root = repository_root / "docs"
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    if not workbook_path.is_file():
        raise RuntimeError("anonymized GL is unavailable")
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    repository = SqlAlchemyRasAuditRepository(
        sessionmaker(bind=engine, expire_on_commit=False)
    )
    attestor = RasFactContextAttestor(
        "report-v2-validation-key-with-at-least-32-bytes"
    )
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=docs_root),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            docs_root / "reference/ras-ledger-account-mapping.organization.csv"
        ),
        ras_fact_context_attestor=attestor,
        ras_audit_repository=repository,
    )
    token = attestor.issue(
        message="Audit RAS anonymise rapport v2.",
        extraction=RasFactExtraction((), (), ()),
        session_id="report-v2-validation-session",
        file_id="report-v2-validation-file",
    )
    batch = executor.execute(
        ToolCall(
            name="run_ras_audit_batch",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
            },
        ),
        ras_fact_context_token=token,
    )
    if not batch.ok:
        raise RuntimeError(batch.error_code or "RAS batch failed")
    audit_id = batch.output.get("audit_id")
    if not isinstance(audit_id, str):
        raise RuntimeError("RAS audit id is unavailable")
    service = PersistedRasAuditReportService(repository)
    report = service.generate(audit_id)
    repeated = service.generate(audit_id)
    if repeated.report_id != report.report_id:
        raise RuntimeError("RAS report is not reproducible")
    generator = RasAuditReportGenerator()
    exports = {
        "json": generator.to_json(report).encode("utf-8"),
        "csv": generator.to_csv(report).encode("utf-8"),
        "xlsx": render_ras_report_xlsx(report),
        "pdf": render_ras_report_pdf(report),
    }
    workbook = load_workbook(BytesIO(exports["xlsx"]), read_only=True)
    if tuple(workbook.sheetnames) != EXPECTED_SHEETS:
        raise RuntimeError("invalid RAS Excel workbook sheets")
    if not PdfReader(BytesIO(exports["pdf"])).pages:
        raise RuntimeError("invalid RAS PDF summary")
    public_text = exports["json"] + exports["csv"]
    if any(
        forbidden in public_text
        for forbidden in (b"candidate_id", b"rule_id", b"action_code")
    ):
        raise RuntimeError("technical identifiers leaked into a public export")
    sizes = ", ".join(f"{name}={len(content)}" for name, content in exports.items())
    print(
        f"report_v2_ok contract={report.contract_version} "
        f"candidates={report.case_count} {sizes}"
    )


if __name__ == "__main__":
    main()
