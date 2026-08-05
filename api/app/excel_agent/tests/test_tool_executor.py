from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from app.database import Base, create_database_engine, create_session_factory
from app.excel_agent.excel_tools import ExcelAgentTools
from app.excel_agent.tests.fixtures import (
    write_large_ras_audit_grand_livre,
    write_minified_grand_livre,
    write_ras_audit_grand_livre,
)
from app.excel_agent.tool_executor import ExcelToolExecutor
from app.excel_agent.tool_registry import create_excel_tool_registry
from app.excel_agent.tool_validator import (
    InvalidToolArgumentsError,
    UnknownToolError,
)
from app.llm.domain import ToolCall
from app.ras_audit.account_mapping import load_ras_ledger_account_mappings
from app.ras_audit.accounting_assessment import (
    load_ras_accounting_assessment_policies,
)
from app.ras_audit.candidate_signals import load_ras_candidate_signals
from app.ras_audit.fact_context import (
    RasExplicitFactExtractor,
    RasExtractedFact,
    RasFactContextAttestor,
    RasFactExtraction,
    load_ras_user_fact_patterns,
)
from app.ras_audit.legal_rules import load_ras_legal_rules
from app.ras_audit.persistence import SqlAlchemyRasAuditRepository
from app.ras_audit.rule_resolution import RasFactSource, RasLegalFact
from app.ras_audit.semantic_classifier import (
    RasTransactionSemanticClassifier,
    load_ras_semantic_policy,
)
from app.ras_audit.tax_event_rules import RasTaxEventRule, load_ras_tax_event_rules
from app.ras_audit.tax_rag_query import TaxRagQueryService
from app.ras_audit.theoretical_calculation import load_ras_calculation_parameters


class UniformSemanticTestProvider:
    def embed_text(self, text: str) -> tuple[float, ...]:
        return (1.0, 0.0)

    def embed_texts(self, texts: tuple[str, ...]) -> tuple[tuple[float, ...], ...]:
        return tuple(self.embed_text(text) for text in texts)


FACT_ATTESTOR = RasFactContextAttestor("test-signing-key-with-at-least-32-bytes")


def _fact_context(
    facts: dict[str, str],
    *,
    session_id: str = "test-session",
    file_id: str = "test-file",
) -> str:
    extraction = RasFactExtraction(
        facts=tuple(
            RasExtractedFact(
                fact=RasLegalFact(
                    name=name,
                    value=value,
                    source=RasFactSource.USER,
                    evidence_reference=f"user_message:test:fact:{name}",
                ),
                start=index,
                end=index + 1,
            )
            for index, (name, value) in enumerate(facts.items())
        ),
        conflicting_fact_names=(),
        pattern_versions=("test-v1",),
    )
    return FACT_ATTESTOR.issue(
        message="test facts",
        extraction=extraction,
        session_id=session_id,
        file_id=file_id,
    )


def test_validator_accepts_known_tool_with_valid_arguments(tmp_path: Path) -> None:
    workbook_path = write_minified_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.validate(
        ToolCall(
            name="profile_sheet",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Grand Livre",
            },
        ),
    )

    assert result.name == "profile_sheet"
    assert result.arguments["sheet_name"] == "Grand Livre"


def test_validator_rejects_unknown_tool(tmp_path: Path) -> None:
    executor = _create_executor(tmp_path)

    with pytest.raises(UnknownToolError, match="unknown tool"):
        executor.validate(ToolCall(name="delete_file", arguments={}))


def test_validator_rejects_missing_required_argument(tmp_path: Path) -> None:
    executor = _create_executor(tmp_path)

    with pytest.raises(InvalidToolArgumentsError, match="sheet_name"):
        executor.validate(
            ToolCall(
                name="profile_sheet",
                arguments={"file_path": "grand_livre_minifie.xlsx"},
            ),
        )


def test_validator_rejects_invalid_argument_type(tmp_path: Path) -> None:
    executor = _create_executor(tmp_path)

    with pytest.raises(InvalidToolArgumentsError, match="file_path"):
        executor.validate(
            ToolCall(
                name="list_sheets",
                arguments={"file_path": 123},
            ),
        )


@pytest.mark.parametrize(
    ("tool_call", "expected_message"),
    (
        (
            ToolCall(
                name="normalize_gl",
                arguments={
                    "file_path": "ledger.xlsx",
                    "sheet_name": "GL",
                    "column_mapping": [],
                },
            ),
            "column_mapping",
        ),
        (
            ToolCall(
                name="query_tax_rag",
                arguments={"query": "RAS", "limit": True},
            ),
            "limit",
        ),
        (
            ToolCall(
                name="query_tax_rag",
                arguments={"query": "RAS", "limit": 6},
            ),
            "limit",
        ),
    ),
)
def test_validator_enforces_non_string_types_and_numeric_bounds(
    tmp_path: Path,
    tool_call: ToolCall,
    expected_message: str,
) -> None:
    executor = _create_executor(tmp_path)

    with pytest.raises(InvalidToolArgumentsError, match=expected_message):
        executor.validate(tool_call)


def test_executor_rejects_file_outside_allowed_root(tmp_path: Path) -> None:
    allowed_root = tmp_path / "allowed"
    forbidden_root = tmp_path / "forbidden"
    allowed_root.mkdir()
    forbidden_root.mkdir()
    workbook_path = write_minified_grand_livre(forbidden_root)
    executor = _create_executor(allowed_root)

    result = executor.execute(
        ToolCall(name="list_sheets", arguments={"file_path": str(workbook_path)}),
    )

    assert result.ok is False
    assert result.error_code == "unsafe_excel_path"
    assert result.output == {}
    assert str(workbook_path) not in str(result.error_message)
    assert str(forbidden_root) not in str(result.error_message)


def test_executor_rejects_missing_file(tmp_path: Path) -> None:
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="list_sheets",
            arguments={"file_path": str(tmp_path / "missing.xlsx")},
        ),
    )

    assert result.ok is False
    assert result.error_code == "excel_file_read"
    assert str(tmp_path) not in str(result.error_message)


def test_executor_rejects_invalid_excel_file(tmp_path: Path) -> None:
    invalid_workbook = tmp_path / "invalid.xlsx"
    invalid_workbook.write_text("not a real workbook", encoding="utf-8")
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="list_sheets",
            arguments={"file_path": str(invalid_workbook)},
        ),
    )

    assert result.ok is False
    assert result.error_code == "excel_file_read"


def test_executor_profiles_minified_grand_livre_without_cell_values(
    tmp_path: Path,
) -> None:
    workbook_path = write_minified_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="profile_sheet",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Grand Livre",
            },
        ),
    )

    assert result.ok is True
    assert result.error_code is None
    assert result.output["sheet_name"] == "Grand Livre"
    assert result.output["row_count"] == 4
    assert result.output["column_count"] == 5
    assert result.output["columns"][0]["name"] == "Compte"
    serialized_output = repr(result.output)
    assert "Achat fournitures" not in serialized_output
    assert "Compte a analyser" not in serialized_output


def test_executor_analyzes_minified_grand_livre_without_cell_values(
    tmp_path: Path,
) -> None:
    workbook_path = write_minified_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="analyze_ledger",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Grand Livre",
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Grand Livre"
    assert result.output["row_count"] == 4
    assert result.output["schema"]["is_valid"] is True
    assert result.output["schema"]["missing_required_columns"] == []
    assert result.output["columns"][0]["name"] == "Compte"
    serialized_output = repr(result.output)
    assert "Achat fournitures" not in serialized_output
    assert "Compte a analyser" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_classifies_anonymized_ledger_schema_without_cell_values() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="classify_ledger_schema",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["schema"]["is_usable"] is True
    assert result.output["schema"]["requires_confirmation"] is False
    mappings = {
        mapping["canonical_field"]: mapping
        for mapping in result.output["schema"]["mappings"]
    }
    assert mappings["account"]["source_column"] == "Compte"
    assert mappings["amount"]["source_column"] == "Montant devise document"
    assert mappings["tax_code"]["source_column"] == "Code TVA"
    assert mappings["document_type"]["source_column"] == "Type de pièce"
    serialized_output = repr(result.output)
    assert "SAP" not in serialized_output
    assert "Achat" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_analyzes_anonymized_ledger_with_canonical_schema() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="analyze_ledger",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["row_count"] == 2500
    assert result.output["schema"]["is_valid"] is True
    assert result.output["schema"]["canonical_schema"]["is_usable"] is True
    assert result.output["schema"]["canonical_schema"]["fields"]["account"] == "Compte"
    assert (
        result.output["schema"]["canonical_schema"]["fields"]["amount"]
        == "Montant devise document"
    )
    assert result.output["schema"]["canonical_schema"]["fields"]["text"] == "Texte"
    serialized_output = repr(result.output)
    assert "Achat" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_aggregates_anonymized_ledger_without_cell_values() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="aggregate_ledger",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "group_by": ["account", "tax_code"],
                "limit": 5,
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["row_count"] == 2500
    assert result.output["amount_field"] == "Montant devise document"
    by_account = result.output["aggregations"]["account"]
    assert by_account["total_groups"] == 13
    account_group = by_account["groups"][0]
    assert account_group["key"] == "61365000"
    assert account_group["entry_count"] == 227
    assert account_group["amount_sum"] == 61173000.0
    assert account_group["balance"] == 61173000.0
    assert account_group["currency"] == "XOF"
    assert account_group["used_entry_count"] == 227
    assert account_group["excluded_entry_count"] == 0
    by_tax_code = result.output["aggregations"]["tax_code"]
    assert by_tax_code["total_groups"] == 2
    tax_group = by_tax_code["groups"][0]
    assert tax_group["key"] == "V1"
    assert tax_group["entry_count"] == 749
    assert tax_group["amount_sum"] == 87676740.0
    assert tax_group["balance"] == 87676740.0
    assert tax_group["currency"] == "XOF"
    serialized_output = repr(result.output)
    assert "Achat" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_queries_ledger_entries_with_pagination_and_allowed_columns() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="query_ledger_entries",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "filters": {"account": "44380002"},
                "page": 1,
                "page_size": 3,
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["total_matches"] == 701
    assert result.output["page"] == 1
    assert result.output["page_size"] == 3
    assert len(result.output["entries"]) == 3
    first_entry = result.output["entries"][0]
    assert first_entry["account"] == "44380002"
    assert "amount" in first_entry
    assert "tax_code" in first_entry
    assert "text" not in first_entry
    serialized_output = repr(result.output)
    assert "Achat" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_queries_real_anonymized_account_with_stable_payload() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="query_ledger_entries",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "filters": {"account": "44585100"},
                "page": 1,
                "page_size": 20,
            },
        ),
    )

    assert result.ok is True
    assert result.output["total_matches"] == 203
    assert result.output["page"] == 1
    assert result.output["page_size"] == 20
    assert result.output["filters"] == {"account": "44585100"}
    assert "account" in result.output["returned_columns"]
    assert len(result.output["entries"]) == 20
    assert result.output["message"] == "Écritures trouvées pour les filtres fournis."


def test_executor_queries_real_anonymized_account_and_period() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="query_ledger_entries",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "filters": {"account": "44585100", "period": "12"},
                "page": 1,
                "page_size": 20,
            },
        ),
    )

    assert result.ok is True
    assert result.output["total_matches"] == 15
    assert result.output["filters"] == {"account": "44585100", "period": "12"}


def test_executor_rejects_invalid_query_filter() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="query_ledger_entries",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "filters": {"unknown": "value"},
            },
        ),
    )

    assert result.ok is False
    assert result.error_code == "invalid_filter"


def test_executor_calculates_ledger_metrics_without_cell_values() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="calculate_ledger_metrics",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "filters": {"account": "44380002"},
                "metrics": ["sum", "count", "average", "min", "max"],
                "top_by": "account",
                "top_limit": 3,
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["total_matches"] == 701
    assert result.output["amount_field"] == "Montant devise document"
    assert result.output["metrics"] == {
        "sum": -124078785.0,
        "count": 701,
        "average": -177002.55,
        "min": -565000.0,
        "max": 346000.0,
    }
    assert result.output["top"]["canonical_field"] == "account"
    top_group = result.output["top"]["groups"][0]
    assert top_group["key"] == "44380002"
    assert top_group["entry_count"] == 701
    assert top_group["amount_sum"] == -124078785.0
    assert top_group["balance"] == -124078785.0
    assert top_group["currency"] == "XOF"
    serialized_output = repr(result.output)
    assert "Achat" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_detects_data_quality_issues_without_cell_values() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="detect_data_quality_issues",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["row_count"] == 2500
    issue_types = {issue["issue_type"] for issue in result.output["issues"]}
    assert "empty_column" in issue_types
    assert "missing_counterparty" in issue_types
    missing_counterparty = next(
        issue
        for issue in result.output["issues"]
        if issue["issue_type"] == "missing_counterparty"
    )
    assert missing_counterparty["affected_count"] == 661
    assert missing_counterparty["severity"] == "warning"
    serialized_output = repr(result.output)
    assert "Achat" not in serialized_output
    assert "601000" not in serialized_output


def test_executor_detects_tax_candidates_from_reference_rules() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = _create_executor(docs_root)

    result = executor.execute(
        ToolCall(
            name="detect_tax_candidates",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "Sheet1",
                "limit": 5,
            },
        ),
    )

    assert result.ok is True
    assert result.output["sheet_name"] == "Sheet1"
    assert result.output["row_count"] == 2500
    assert result.output["decision_status"] == "review_required"
    categories = {
        candidate["category"]: candidate for candidate in result.output["candidates"]
    }
    assert categories["resident_services"]["entry_count"] == 227
    assert categories["resident_services"]["amount_sum"] == 61173000.0
    assert categories["real_estate_charges"]["entry_count"] == 129
    assert categories["non_resident_services"]["entry_count"] == 87
    serialized_output = repr(result.output)
    assert "HONORAIRES" not in serialized_output
    assert "LOYER" not in serialized_output


def test_executor_normalizes_gl_without_returning_cell_values(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="normalize_gl",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "column_mapping": _ras_audit_column_mapping(),
            },
        ),
    )

    assert result.ok is True
    assert result.output["row_count"] == 3
    assert result.output["normalized_count"] == 3
    assert result.output["rejected_count"] == 0
    assert result.output["issue_counts"] == {}
    readiness = {item["capability"]: item for item in result.output["readiness"]}
    counterpart = readiness["counterpart_reconciliation"]
    assert counterpart["status"] == "blocked"
    assert "source_scope_not_confirmed_complete" in counterpart["blockers"]
    serialized_output = repr(result.output)
    assert "0632100" not in serialized_output
    assert "SYN-TIERS-001" not in serialized_output
    assert "Honoraires synthetiques" not in serialized_output


def test_executor_rejects_duplicate_normalization_mapping(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="normalize_gl",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "column_mapping": {
                    "account_number": "Compte",
                    "partner_id": "Compte",
                },
            },
        ),
    )

    assert result.ok is False
    assert result.error_code == "invalid_column_mapping"


def test_executor_normalizes_gl_with_versioned_automatic_aliases(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="normalize_gl",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["normalized_count"] == 3
    assert result.output["mapped_fields"] == _ras_audit_column_mapping()


def test_executor_resolves_partner_columns_and_flags_row_conflicts(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["GL"]
    sheet["I1"] = "Client"
    sheet["N1"] = "Fournisseur"
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row=row_number, column=14).value = f"F-{row_number}"
    workbook.save(workbook_path)

    result = _create_executor(tmp_path).execute(
        ToolCall(
            name="normalize_gl",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["issue_counts"]["conflicting_vendor_customer"] == 3
    assert result.output["mapped_fields"]["vendor_id"] == "Fournisseur"
    assert result.output["mapped_fields"]["customer_id"] == "Client"


def test_executor_reads_only_the_explicitly_selected_sheet(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    workbook = load_workbook(workbook_path)
    archive = workbook.copy_worksheet(workbook["GL"])
    archive.title = "GL archive"
    archive["M2"] = "FCFA"
    workbook.save(workbook_path)

    result = _create_executor(tmp_path).execute(
        ToolCall(
            name="find_ras_counterpart",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["source_scope_complete"] is True
    assert result.output["source_scope_blockers"] == []


def test_executor_assesses_gl_readiness_without_returning_hash_or_cells(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="assess_gl_readiness",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["row_count"] == 3
    assert "content_sha256" not in result.output
    assert "normalized_count" not in result.output
    readiness = {item["capability"]: item for item in result.output["readiness"]}
    assert readiness["candidate_detection"]["status"] == "ready"
    counterpart = readiness["counterpart_reconciliation"]
    assert counterpart["status"] == "blocked"
    assert "source_scope_not_confirmed_complete" in counterpart["blockers"]
    serialized_output = repr(result.output)
    assert "SYN-TIERS-001" not in serialized_output
    assert "Honoraires synthetiques" not in serialized_output


def test_readiness_is_blocked_without_configured_ras_account_mapping(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
    )

    result = executor.execute(
        ToolCall(
            name="assess_gl_readiness",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    readiness = {item["capability"]: item for item in result.output["readiness"]}
    reconciliation = readiness["counterpart_reconciliation"]
    assert reconciliation["status"] == "blocked"
    assert "ras_account_mapping_unavailable" in reconciliation["blockers"]


def test_llm_cannot_claim_ras_account_mapping_is_available(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="assess_gl_readiness",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "ras_account_mapping_available": True,
            },
        ),
    )

    assert result.ok is False
    assert result.error_code == "invalid_tool_call"


def test_executor_reconstructs_accounting_entries_without_returning_cells(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="reconstruct_accounting_entry",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["row_count"] == 3
    assert result.output["entry_count"] == 1
    assert result.output["grouped_line_count"] == 3
    assert result.output["balanced_count"] == 1
    assert result.output["unbalanced_count"] == 0
    assert result.output["ungrouped_line_count"] == 0
    assert result.output["currencies"] == ["XOF"]
    serialized_output = repr(result.output)
    assert "000042" not in serialized_output
    assert "0632100" not in serialized_output
    assert "SYN-TIERS-001" not in serialized_output


def test_executor_finds_ras_counterpart_with_structured_safe_summary(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="find_ras_counterpart",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "related_window_days": 31,
            },
        ),
    )

    assert result.ok is True
    assert result.output["candidate_piece_count"] == 1
    assert result.output["status_counts"]["found_in_same_entry"] == 1
    assert result.output["confirmed_amounts_by_currency"] == {"XOF": "5000"}
    assert result.output["potential_related_amounts_by_currency"] == {}
    assert result.output["source_scope_complete"] is True
    serialized_output = repr(result.output)
    assert "000042" not in serialized_output
    assert "0447100" not in serialized_output
    assert "SYN-TIERS-001" not in serialized_output


def test_executor_explains_incomplete_scope_for_invalid_currency(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    workbook = load_workbook(workbook_path)
    workbook["GL"]["M2"] = "FCFA"
    workbook.save(workbook_path)

    result = _create_executor(tmp_path).execute(
        ToolCall(
            name="find_ras_counterpart",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["source_scope_complete"] is False
    assert "missing_or_invalid_currencies" in result.output[
        "source_scope_blockers"
    ]
    assert result.output["source_scope_policy_version"] == (
        "ras-uploaded-sheet-scope-v4"
    )
    assert result.output["issue_counts"]["invalid_currency"] == 1


def test_executor_blocks_conflicting_document_currencies(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    workbook = load_workbook(workbook_path)
    workbook["GL"]["M3"] = "USD"
    workbook.save(workbook_path)

    result = _create_executor(tmp_path).execute(
        ToolCall(
            name="find_ras_counterpart",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["source_scope_complete"] is False
    assert "conflicting_document_currencies" in result.output[
        "source_scope_blockers"
    ]
    assert result.output["status_counts"]["found_in_same_entry"] == 1


def test_executor_detects_ras_candidates_with_structured_safe_summary(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="detect_ras_candidates",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["evaluated_piece_count"] == 1
    assert result.output["candidate_piece_count"] == 1
    assert result.output["status_counts"] == {"candidate_account_and_text": 1}
    assert result.output["candidate_amounts_by_currency"] == {"XOF": "100000"}
    assert result.output["ras_review"] == {
        "periods": [
            {
                "period": "1",
                "evaluated_entry_count": 1,
                "candidate_entry_count": 1,
                "candidate_amount": "100000",
                "currency": "XOF",
                "candidate_rate": 1.0,
            },
        ],
    }
    assert result.output["decision_status"] == "review_only_no_tax_conclusion"
    serialized_output = repr(result.output)
    assert "000042" not in serialized_output
    assert "0632100" not in serialized_output
    assert "SYN-TIERS-001" not in serialized_output
    assert "Honoraires synthetiques" not in serialized_output


def test_real_anonymized_gl_resolves_vendor_customer_and_organization_mapping() -> None:
    docs_root = _docs_root()
    workbook_path = docs_root / "GL_anonymise_2500.xlsx"
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=docs_root),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            docs_root / "reference/ras-ledger-account-mapping.organization.csv"
        ),
    )

    candidates = executor.execute(
        ToolCall(
            name="detect_ras_candidates",
            arguments={"file_path": str(workbook_path), "sheet_name": "Sheet1"},
        )
    )
    counterparts = executor.execute(
        ToolCall(
            name="find_ras_counterpart",
            arguments={"file_path": str(workbook_path), "sheet_name": "Sheet1"},
        )
    )

    assert candidates.ok is True
    assert candidates.output["candidate_piece_count"] == 611
    assert candidates.output["rejected_row_count"] == 0
    assert counterparts.ok is True
    assert counterparts.output["status_counts"]["found_in_same_entry"] == 59
    assert counterparts.output["source_scope_complete"] is False
    assert counterparts.output["source_scope_blockers"] == [
        "journal_is_document_type_proxy",
        "missing_company_scope",
        "posting_date_is_document_date_proxy",
    ]
    assert counterparts.output["status_counts"]["not_found_in_scope"] == 0


def test_explicit_default_company_code_can_attest_a_single_company_extract(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["GL"]
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row=row_number, column=1).value = None
    workbook.save(workbook_path)
    workbook.close()
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            Path("app/ras_audit/tests/fixtures/account-mapping/valid.csv")
        ),
        default_company_code="SYN-CO",
    )

    result = executor.execute(
        ToolCall(
            name="find_ras_counterpart",
            arguments={"file_path": str(workbook_path), "sheet_name": "GL"},
        )
    )

    assert result.ok is True
    assert result.output["source_scope_complete"] is True
    assert result.output["source_scope_blockers"] == []
    assert result.output["issue_counts"][
        "company_code_from_organization_config"
    ] == 3


def test_executor_classifies_semantics_without_returning_labels(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            Path("app/ras_audit/tests/fixtures/account-mapping/valid.csv"),
        ),
        ras_semantic_classifier=RasTransactionSemanticClassifier(
            embedding_provider=UniformSemanticTestProvider(),
            provider_name="semantic-test",
            model_name="synthetic-model-v1",
            signals=load_ras_candidate_signals(
                Path("../docs/reference/ras-candidate-signals.csv")
            ),
            policy=load_ras_semantic_policy(
                Path("../docs/reference/ras-semantic-classification-policy.csv")
            ),
        ),
    )

    result = executor.execute(
        ToolCall(
            name="classify_transaction_semantics",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is True
    assert result.output["classified_piece_count"] == 1
    assert result.output["status_counts"]["semantic_ambiguous"] == 1
    assert result.output["model_name"] == "synthetic-model-v1"
    assert result.output["policy_version"] == "1.0.0"
    assert result.output["decision_status"] == "semantic_suggestion_for_review_only"
    serialized_output = repr(result.output)
    assert "Honoraires synthetiques" not in serialized_output
    assert "SYN-TIERS-001" not in serialized_output


def test_executor_refuses_semantic_tool_when_model_is_disabled(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    result = _create_executor(tmp_path).execute(
        ToolCall(
            name="classify_transaction_semantics",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
    )

    assert result.ok is False
    assert result.error_code == "ras_semantic_classifier_unavailable"


def test_executor_resolves_legal_rule_without_exposing_fact_values(
    tmp_path: Path,
) -> None:
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_legal_rules=load_ras_legal_rules(
            Path("../docs/reference/bf-ras-legal-rules.csv"),
            repository_root=Path(".."),
        ),
        ras_tax_event_rules=_tax_event_rules(),
        ras_fact_context_attestor=FACT_ATTESTOR,
    )
    result = executor.execute(
        ToolCall(
            name="resolve_applicable_ras_rule",
            arguments={
                "transaction_date": "2026-04-10",
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "regime": "resident",
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "987654321",
                "currency": "XOF",
                "payment_status": "paid",
                "payment_date": "2026-04-10",
            }
        ),
    )

    assert result.ok is True
    assert result.output["status"] == "resolved_provisional"
    assert result.output["rule_id"] == "resident_standard_2026"
    assert result.output["rate_percent"] == "5"
    attestation = result.output["fact_attestation"]
    assert attestation["message_sha256"]
    assert "residence_status" in attestation["fact_names"]
    assert attestation["conflicting_fact_names"] == []
    assert "token" not in repr(attestation).lower()
    assert "100000" not in repr(attestation)
    assert len(result.output["sources"]) == 3
    assert "987654321" not in repr(result.output)


def test_executor_resolves_rule_from_real_explicit_user_message(
    tmp_path: Path,
) -> None:
    message = (
        "Le prestataire résident est un prestataire immatriculé à l'IFU. "
        "Il s'agit d'une prestation de services, avec un payeur éligible à la "
        "retenue sur prestations et un service utilisé au Burkina Faso. "
        "Aucune exonération applicable. L'assiette fiscale est de 100 000 XOF."
    )
    message += " Le paiement effectu\u00e9 le 2026-04-10."
    extraction = RasExplicitFactExtractor(
        load_ras_user_fact_patterns(
            Path("../docs/reference/ras-user-fact-patterns.csv")
        )
    ).extract(message)
    token = FACT_ATTESTOR.issue(
        message=message,
        extraction=extraction,
        session_id="session-1",
        file_id="file-1",
    )
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_legal_rules=load_ras_legal_rules(
            Path("../docs/reference/bf-ras-legal-rules.csv"),
            repository_root=Path(".."),
        ),
        ras_tax_event_rules=_tax_event_rules(),
        ras_fact_context_attestor=FACT_ATTESTOR,
    )

    result = executor.execute(
        ToolCall(
            name="resolve_applicable_ras_rule",
            arguments={"transaction_date": "2026-04-10"},
        ),
        ras_fact_context_token=token,
    )

    assert result.ok is True
    assert result.output["status"] == "resolved_provisional"
    assert result.output["rule_id"] == "resident_standard_2026"
    assert result.output["rate_percent"] == "5"


def test_executor_rejects_free_form_legal_facts(tmp_path: Path) -> None:
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_legal_rules=load_ras_legal_rules(
            Path("../docs/reference/bf-ras-legal-rules.csv"),
            repository_root=Path(".."),
        ),
        ras_tax_event_rules=_tax_event_rules(),
    )
    result = executor.execute(
        ToolCall(
            name="resolve_applicable_ras_rule",
            arguments={
                "transaction_date": "2026-04-10",
                "facts": {"llm_guessed_rate": "5"},
            },
        )
    )

    assert result.ok is False
    assert result.error_code == "invalid_tool_call"


def test_executor_refuses_report_without_persisted_audit_repository(
    tmp_path: Path,
) -> None:
    result = _create_executor(tmp_path).execute(
        ToolCall(
            name="generate_ras_audit_report",
            arguments={"audit_id": "audit-1"},
        )
    )

    assert result.ok is False
    assert result.error_code == "ras_audit_repository_unavailable"


def test_executor_returns_structured_tax_rag_citations(tmp_path: Path) -> None:
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        tax_rag_query_service=TaxRagQueryService(Path("../docs/source-corpus/fiscal")),
    )

    result = executor.execute(
        ToolCall(
            name="query_tax_rag",
            arguments={
                "query": "retenue à la source prestataire résident",
                "as_of_date": "2026-04-10",
                "limit": 2,
            },
        )
    )

    assert result.ok is True
    assert result.output["citations"]
    assert result.output["decision_status"] == "retrieval_only_no_tax_decision"


def test_executor_calculates_theoretical_ras_with_full_trace(
    tmp_path: Path,
) -> None:
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_legal_rules=load_ras_legal_rules(
            Path("../docs/reference/bf-ras-legal-rules.csv"),
            repository_root=Path(".."),
        ),
        ras_tax_event_rules=_tax_event_rules(),
        ras_calculation_parameters=load_ras_calculation_parameters(
            Path("../docs/reference/bf-ras-calculation-parameters.csv"),
            repository_root=Path(".."),
        ),
        ras_fact_context_attestor=FACT_ATTESTOR,
    )
    result = executor.execute(
        ToolCall(
            name="calculate_theoretical_ras",
            arguments={
                "transaction_date": "2026-04-10",
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "regime": "resident",
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000.50",
                "currency": "XOF",
                "payment_status": "paid",
                "payment_date": "2026-04-10",
            }
        ),
    )

    assert result.ok is True
    assert result.output["legal_resolution_status"] == "resolved_provisional"
    assert result.output["calculation_status"] == "calculated_provisional"
    assert result.output["base_amount"] == "100000.50"
    assert result.output["expected_amount"] == "5000.025"
    assert result.output["currency"] == "XOF"
    assert result.output["rounding_policy"] == (
        "none_exact_decimal_source_rounding_unavailable"
    )
    assert len(result.output["legal_sources"]) == 3
    assert {source["source_kind"] for source in result.output["legal_sources"]} == {
        "scope",
        "rate",
        "tax_event",
    }


def test_executor_refuses_unresolved_partial_payment_allocation(
    tmp_path: Path,
) -> None:
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_legal_rules=load_ras_legal_rules(
            Path("../docs/reference/bf-ras-legal-rules.csv"),
            repository_root=Path(".."),
        ),
        ras_tax_event_rules=_tax_event_rules(),
        ras_calculation_parameters=load_ras_calculation_parameters(
            Path("../docs/reference/bf-ras-calculation-parameters.csv"),
            repository_root=Path(".."),
        ),
        ras_fact_context_attestor=FACT_ATTESTOR,
    )
    result = executor.execute(
        ToolCall(
            name="calculate_theoretical_ras",
            arguments={"transaction_date": "2026-04-10"},
        ),
        ras_fact_context_token=_fact_context(
            {
                "regime": "resident",
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000",
                "payment_amount": "40000",
                "currency": "XOF",
                "payment_status": "paid",
                "payment_date": "2026-04-10",
            }
        ),
    )

    assert result.ok is True
    assert result.output["legal_resolution_status"] == "resolved_provisional"
    assert result.output["calculation_status"] == "not_calculable"
    assert result.output["expected_amount"] is None
    assert result.output["reason"] == (
        "partial_payment_tax_base_allocation_unresolved"
    )


def test_executor_assesses_selected_ras_entry_end_to_end(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    executor = _create_ras_assessment_executor(tmp_path)

    result = executor.execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "accounting_entry": {
                    "company_code": "SYN-CO",
                    "fiscal_year": 2025,
                    "journal": "ACH",
                    "document_number": "000042",
                },
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "regime": "resident",
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000",
                "payment_status": "paid",
                "payment_date": "2025-01-10",
            }
        ),
    )

    assert result.ok is True
    assert result.output["status"] == "provisional_reconciled"
    assert result.output["expected_amount"] == "5000"
    assert result.output["recorded_amount"] == "5000"
    assert result.output["difference"] == "0"
    assert result.output["currency"] == "XOF"
    assert result.output["basis_is_complete"] is True
    serialized = repr(result.output)
    assert "SYN-CO" not in serialized
    assert "000042" not in serialized
    assert "SYN-TIERS-001" not in serialized


def test_executor_dates_rule_from_attested_payment_not_expense_posting(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    result = _create_ras_assessment_executor(tmp_path).execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "accounting_entry": {
                    "company_code": "SYN-CO",
                    "fiscal_year": 2025,
                    "journal": "ACH",
                    "document_number": "000042",
                },
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000",
                "payment_status": "paid",
                "payment_date": "2026-01-10",
            }
        ),
    )

    assert result.ok is True
    assert result.output["rule_id"] == "resident_standard_2026"
    assert result.output["rule_version"] == "v2026"


def test_executor_does_not_calculate_unpaid_or_undated_charge(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    result = _create_ras_assessment_executor(tmp_path).execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "accounting_entry": {
                    "company_code": "SYN-CO",
                    "fiscal_year": 2025,
                    "journal": "ACH",
                    "document_number": "000042",
                },
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000",
            }
        ),
    )

    assert result.ok is True
    assert result.output["legal_resolution_status"] == "missing_facts"
    assert result.output["calculation_status"] == "not_calculable"
    assert result.output["status"] == "indeterminate"


def test_executor_refuses_single_rate_for_multi_category_piece(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    workbook = load_workbook(workbook_path)
    workbook["GL"]["J2"] = "Honoraires loyer synthetiques"
    workbook.save(workbook_path)
    result = _create_ras_assessment_executor(tmp_path).execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "accounting_entry": {
                    "company_code": "SYN-CO",
                    "fiscal_year": 2025,
                    "journal": "ACH",
                    "document_number": "000042",
                },
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000",
            }
        ),
    )

    assert result.ok is False
    assert result.error_code == "ras_candidate_requires_category_split"


def test_executor_persists_assessment_then_generates_report(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    repository = SqlAlchemyRasAuditRepository(_session_factory(tmp_path))
    executor = _create_ras_assessment_executor(tmp_path, repository)
    fact_token = _fact_context(
        {
            "regime": "resident",
            "operation_type": "service_any",
            "residence_status": "resident",
            "ifu_status": "registered",
            "payer_type": "eligible_service_payer",
            "service_use_location": "BF",
            "exemption_status": "not_exempt",
            "tax_base_amount": "100000",
            "payment_status": "paid",
            "payment_date": "2025-01-10",
        }
    )

    assessment = executor.execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "accounting_entry": {
                    "company_code": "SYN-CO",
                    "fiscal_year": 2025,
                    "journal": "ACH",
                    "document_number": "000042",
                },
            },
        ),
        ras_fact_context_token=fact_token,
    )
    report = executor.execute(
        ToolCall(
            name="generate_ras_audit_report",
            arguments={"audit_id": assessment.output["audit_id"]},
        ),
        ras_fact_context_token=fact_token,
    )

    assert assessment.ok is True
    assert report.ok is True
    assert report.output["case_count"] == 1
    assert report.output["amount_summaries"][0]["expected_amount"] == "5000"
    persisted = repository.get(str(assessment.output["audit_id"]))
    assert persisted is not None
    assert "tax-event:v1" in persisted.reference_versions

    foreign_report = executor.execute(
        ToolCall(
            name="generate_ras_audit_report",
            arguments={"audit_id": assessment.output["audit_id"]},
        ),
        ras_fact_context_token=_fact_context({}, session_id="other-session"),
    )
    assert foreign_report.ok is False
    assert foreign_report.error_code == "ras_audit_access_denied"


def test_executor_persists_multi_candidate_gl_inventory(tmp_path: Path) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    repository = SqlAlchemyRasAuditRepository(_session_factory(tmp_path))
    executor = _create_ras_assessment_executor(tmp_path, repository)
    batch_token = _fact_context({})

    batch = executor.execute(
        ToolCall(
            name="run_ras_audit_batch",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "max_candidates": 100,
            },
        ),
        ras_fact_context_token=batch_token,
    )
    report = executor.execute(
        ToolCall(
            name="generate_ras_audit_report",
            arguments={"audit_id": batch.output["audit_id"]},
        ),
        ras_fact_context_token=batch_token,
    )

    assert batch.ok is True
    assert batch.output["candidate_count"] >= 1
    assert len(batch.output["review_candidate_ids"]) >= 1
    assert batch.output["remaining_candidate_count"] == 0
    assert batch.output["decision_status"] == (
        "candidate_inventory_pending_legal_facts"
    )
    assert report.ok is True
    assert report.output["case_count"] == batch.output["candidate_count"]
    assert all(
        detail["basis_is_complete"] is False for detail in report.output["details"]
    )

    enriched = executor.execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "base_audit_id": batch.output["audit_id"],
                "candidate_id": batch.output["review_candidate_ids"][0],
            },
        ),
        ras_fact_context_token=_fact_context(
            {
                "regime": "resident",
                "operation_type": "service_any",
                "residence_status": "resident",
                "ifu_status": "registered",
                "payer_type": "eligible_service_payer",
                "service_use_location": "BF",
                "exemption_status": "not_exempt",
                "tax_base_amount": "100000",
                "payment_status": "paid",
                "payment_date": "2025-01-10",
            }
        ),
    )
    derived = repository.get(str(enriched.output["audit_id"]))

    assert enriched.ok is True
    assert derived is not None
    assert derived.parent_audit_id == batch.output["audit_id"]
    assert len(derived.cases) == batch.output["candidate_count"]
    assert any(case.status == "provisional_reconciled" for case in derived.cases)

    unknown_candidate = executor.execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "base_audit_id": batch.output["audit_id"],
                "candidate_id": "f" * 64,
            },
        ),
        ras_fact_context_token=_fact_context({"currency": "XOF"}),
    )
    assert unknown_candidate.ok is False
    assert unknown_candidate.error_code == "ras_candidate_not_found"


def test_executor_processes_large_batch_with_bounded_agent_output(
    tmp_path: Path,
) -> None:
    workbook_path = write_large_ras_audit_grand_livre(
        tmp_path,
        candidate_count=600,
    )
    repository = SqlAlchemyRasAuditRepository(_session_factory(tmp_path))
    result = _create_ras_assessment_executor(tmp_path, repository).execute(
        ToolCall(
            name="run_ras_audit_batch",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
            },
        ),
        ras_fact_context_token=_fact_context({}),
    )

    assert result.ok is True
    assert result.output["candidate_count"] == 600
    assert len(result.output["review_candidate_ids"]) == 20
    assert result.output["remaining_candidate_count"] == 580
    assert result.output["source_scope_complete"] is True
    persisted = repository.get(str(result.output["audit_id"]))
    assert persisted is not None
    assert len(persisted.cases) == 600


def test_executor_enforces_server_candidate_limit(tmp_path: Path) -> None:
    workbook_path = write_large_ras_audit_grand_livre(
        tmp_path,
        candidate_count=2,
    )
    repository = SqlAlchemyRasAuditRepository(_session_factory(tmp_path))
    executor = ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=tmp_path),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            Path("app/ras_audit/tests/fixtures/account-mapping/valid.csv"),
        ),
        ras_fact_context_attestor=FACT_ATTESTOR,
        ras_audit_repository=repository,
        max_ras_batch_candidates=1,
    )

    result = executor.execute(
        ToolCall(
            name="run_ras_audit_batch",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "max_candidates": 20_000,
            },
        ),
        ras_fact_context_token=_fact_context({}),
    )

    assert result.ok is False
    assert result.error_code == "ras_batch_limit_exceeded"


def test_executor_rejects_currency_fact_different_from_selected_entry(
    tmp_path: Path,
) -> None:
    workbook_path = write_ras_audit_grand_livre(tmp_path)
    result = _create_ras_assessment_executor(tmp_path).execute(
        ToolCall(
            name="assess_ras_accounting",
            arguments={
                "file_path": str(workbook_path),
                "sheet_name": "GL",
                "accounting_entry": {
                    "company_code": "SYN-CO",
                    "fiscal_year": 2025,
                    "journal": "ACH",
                    "document_number": "000042",
                },
            },
        ),
        ras_fact_context_token=_fact_context({"currency": "EUR"}),
    )

    assert result.ok is False
    assert result.error_code == "currency_fact_mismatch"


def _create_executor(allowed_root: Path) -> ExcelToolExecutor:
    return ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=allowed_root),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            Path("app/ras_audit/tests/fixtures/account-mapping/valid.csv"),
        ),
    )


def _tax_event_rules() -> tuple[RasTaxEventRule, ...]:
    return load_ras_tax_event_rules(
        Path("../docs/reference/bf-ras-tax-event-rules.csv"),
        repository_root=Path(".."),
    )


def _create_ras_assessment_executor(
    allowed_root: Path,
    repository: SqlAlchemyRasAuditRepository | None = None,
) -> ExcelToolExecutor:
    return ExcelToolExecutor(
        tools=ExcelAgentTools(allowed_root=allowed_root),
        registry=create_excel_tool_registry(),
        ras_ledger_account_mappings=load_ras_ledger_account_mappings(
            Path("app/ras_audit/tests/fixtures/account-mapping/valid.csv"),
        ),
        ras_legal_rules=load_ras_legal_rules(
            Path("../docs/reference/bf-ras-legal-rules.csv"),
            repository_root=Path(".."),
        ),
        ras_tax_event_rules=_tax_event_rules(),
        ras_calculation_parameters=load_ras_calculation_parameters(
            Path("../docs/reference/bf-ras-calculation-parameters.csv"),
            repository_root=Path(".."),
        ),
        ras_accounting_assessment_policies=(
            load_ras_accounting_assessment_policies(
                Path("../docs/reference/ras-accounting-assessment-policy.csv")
            )
        ),
        ras_fact_context_attestor=FACT_ATTESTOR,
        ras_audit_repository=repository,
    )


def _session_factory(tmp_path: Path) -> sessionmaker[Session]:
    engine = create_database_engine(f"sqlite:///{tmp_path / 'executor.db'}")
    Base.metadata.create_all(engine)
    return create_session_factory(engine)


def _docs_root() -> Path:
    docker_docs_root = Path("/workspace/docs")
    if docker_docs_root.is_dir():
        return docker_docs_root
    return Path("../docs")


def _ras_audit_column_mapping() -> dict[str, str]:
    return {
        "company_code": "Societe",
        "fiscal_year": "Exercice",
        "period": "Periode",
        "journal": "Journal",
        "document_number": "Numero piece",
        "line_number": "Numero ligne",
        "posting_date": "Date comptable",
        "account_number": "Compte",
        "partner_id": "Tiers",
        "label": "Libelle",
        "posting_key": "Cle de comptabilisation",
        "amount": "Montant devise document",
        "currency": "Devise du document",
    }
