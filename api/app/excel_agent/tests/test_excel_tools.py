from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.excel_agent.domain import (
    ExcelFileReadError,
    ExcelSheetNotFoundError,
    UnsafeExcelPathError,
    UnsupportedExcelFileError,
)
from app.excel_agent.excel_tools import ExcelAgentTools


def test_list_sheets_returns_workbook_sheet_names(tmp_path: Path) -> None:
    workbook_path = _write_workbook(tmp_path)
    tools = ExcelAgentTools(allowed_root=tmp_path)

    result = tools.list_sheets(workbook_path)

    assert result.sheet_names == ("Grand Livre", "Plan Comptable")
    assert [(sheet.name, sheet.visibility) for sheet in result.sheets] == [
        ("Grand Livre", "visible"),
        ("Plan Comptable", "visible"),
    ]


def test_get_columns_returns_normalized_headers(tmp_path: Path) -> None:
    workbook_path = _write_workbook(tmp_path)
    tools = ExcelAgentTools(allowed_root=tmp_path)

    result = tools.get_columns(workbook_path, sheet_name="Grand Livre")

    assert result.sheet_name == "Grand Livre"
    assert result.columns == ("Compte", "Libelle", "Montant", "Date")


def test_get_columns_preserves_special_headers_and_rejects_duplicates(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "headers.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        pd.DataFrame([[1, 2, 3, 4]]).to_excel(
            writer,
            sheet_name="Écritures & RAS",
            index=False,
            header=["  Montant  ", "", "Débit/Crédit", "Montant"],
        )
    tools = ExcelAgentTools(allowed_root=tmp_path)

    with pytest.raises(ExcelFileReadError, match="duplicate"):
        tools.get_columns(workbook_path, sheet_name="Écritures & RAS")


def test_list_sheets_exposes_hidden_state_and_releases_workbook(
    tmp_path: Path,
) -> None:
    workbook_path = _write_workbook(tmp_path)
    workbook = load_workbook(workbook_path)
    workbook["Plan Comptable"].sheet_state = "hidden"
    workbook.create_sheet("Feuille vide & spéciale")
    workbook.save(workbook_path)
    workbook.close()
    tools = ExcelAgentTools(allowed_root=tmp_path)

    result = tools.list_sheets(workbook_path)

    assert result.sheet_names == (
        "Grand Livre",
        "Plan Comptable",
        "Feuille vide & spéciale",
    )
    assert result.sheets[1].visibility == "hidden"
    released_path = tmp_path / "released.xlsx"
    workbook_path.rename(released_path)
    assert released_path.is_file()


def test_profile_sheet_returns_statistics_without_cell_values(tmp_path: Path) -> None:
    workbook_path = _write_workbook(tmp_path)
    tools = ExcelAgentTools(allowed_root=tmp_path)

    profile = tools.profile_sheet(workbook_path, sheet_name="Grand Livre")

    assert profile.sheet_name == "Grand Livre"
    assert profile.row_count == 3
    assert profile.column_count == 4
    assert [column.name for column in profile.columns] == [
        "Compte",
        "Libelle",
        "Montant",
        "Date",
    ]
    assert profile.columns[0].non_empty_count == 3
    assert profile.columns[1].missing_count == 1
    assert profile.columns[1].detected_type == "text"
    assert profile.columns[2].detected_type == "number"
    assert profile.columns[3].detected_type == "datetime"
    serialized = repr(profile)
    assert "Achat fournitures" not in serialized
    assert "Prestation conseil" not in serialized


def test_read_sheet_rows_reuses_only_an_unchanged_file_within_executor(
    tmp_path: Path,
) -> None:
    workbook_path = _write_workbook(tmp_path)
    tools = ExcelAgentTools(allowed_root=tmp_path)

    first = tools.read_sheet_rows(workbook_path, "Grand Livre")
    second = tools.read_sheet_rows(workbook_path, "Grand Livre")
    assert second is first

    workbook = load_workbook(workbook_path)
    workbook["Grand Livre"].append(["613000", "Nouvelle ligne", 10, None])
    workbook.save(workbook_path)
    workbook.close()

    refreshed = tools.read_sheet_rows(workbook_path, "Grand Livre")
    assert refreshed is not first
    assert refreshed.row_count == first.row_count + 1


def test_tools_reject_files_outside_allowed_root(tmp_path: Path) -> None:
    allowed_root = tmp_path / "allowed"
    forbidden_root = tmp_path / "forbidden"
    allowed_root.mkdir()
    forbidden_root.mkdir()
    workbook_path = _write_workbook(forbidden_root)
    tools = ExcelAgentTools(allowed_root=allowed_root)

    with pytest.raises(UnsafeExcelPathError):
        tools.list_sheets(workbook_path)


def test_tools_accept_files_from_secondary_allowed_root(tmp_path: Path) -> None:
    corpus_root = tmp_path / "corpus"
    session_root = tmp_path / "sessions"
    corpus_root.mkdir()
    session_root.mkdir()
    workbook_path = _write_workbook(session_root)
    tools = ExcelAgentTools(
        allowed_root=corpus_root,
        allowed_roots=(session_root,),
    )

    result = tools.list_sheets(workbook_path)

    assert result.sheet_names == ("Grand Livre", "Plan Comptable")


def test_tools_reject_unsupported_excel_extension(tmp_path: Path) -> None:
    source_path = tmp_path / "source.csv"
    source_path.write_text("Compte,Montant\n601,10\n", encoding="utf-8")
    tools = ExcelAgentTools(allowed_root=tmp_path)

    with pytest.raises(UnsupportedExcelFileError):
        tools.list_sheets(source_path)


def test_get_columns_rejects_unknown_sheet(tmp_path: Path) -> None:
    workbook_path = _write_workbook(tmp_path)
    tools = ExcelAgentTools(allowed_root=tmp_path)

    with pytest.raises(ExcelSheetNotFoundError):
        tools.get_columns(workbook_path, sheet_name="Inconnue")


def _write_workbook(directory: Path) -> Path:
    workbook_path = directory / "harmonizer.xlsx"
    ledger = pd.DataFrame(
        {
            "Compte": ["601000", "604000", "706000"],
            "Libelle": ["Achat fournitures", None, "Prestation conseil"],
            "Montant": [1200.5, None, 900.0],
            "Date": pd.to_datetime(["2026-01-01", "2026-01-02", "2026-01-03"]),
        },
    )
    plan = pd.DataFrame(
        {
            "Compte": ["601000"],
            "Texte descr.cpt gen.": ["Achats stockes"],
        },
    )
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        ledger.to_excel(writer, sheet_name="Grand Livre", index=False)
        plan.to_excel(writer, sheet_name="Plan Comptable", index=False)
    return workbook_path
