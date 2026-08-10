from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.ras_audit.audit_report import (
    RasAuditReport,
    RasAuditReportDetail,
    ras_report_action_label,
    ras_report_certainty_label,
    ras_report_fact_label,
    ras_report_operation_label,
    ras_report_priority_label,
)

RAS_REPORT_EXPORT_MAX_BYTES = 20 * 1024 * 1024

DETAIL_HEADERS = (
    "Période",
    "Pièce",
    "Compte",
    "Nature de l'opération",
    "Fournisseur",
    "Statut",
    "Niveau de conclusion",
    "Priorité",
    "Action recommandée",
    "Constat",
    "Assiette fiscale",
    "Taux (%)",
    "RAS théorique",
    "RAS comptabilisée",
    "Écart comptabilisé - théorique",
    "Devise",
    "Informations manquantes",
    "Anomalies",
    "Références juridiques",
)


def render_ras_report_xlsx(report: RasAuditReport) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Synthese"
    _write_summary(summary, report)
    _write_detail_sheet(
        workbook,
        "Anomalies",
        tuple(
            detail
            for detail in report.details
            if detail.review_priority == "high" or detail.issues
        ),
    )
    _write_detail_sheet(
        workbook,
        "Dossiers incomplets",
        tuple(
            detail
            for detail in report.details
            if detail.certainty.value == "indeterminate"
        ),
    )
    target = BytesIO()
    workbook.save(target)
    return _bounded(target.getvalue())


def render_ras_report_pdf(report: RasAuditReport) -> bytes:
    lines = (
        "RAPPORT DE SYNTHESE - AUDIT RAS",
        f"Généré le : {report.generated_at.isoformat()}",
        f"Dossiers analysés : {report.case_count}",
        f"Cas conclus: {report.executive_summary.concluded_count}",
        f"Cas probables: {report.executive_summary.probable_count}",
        f"Cas indeterminables: {report.executive_summary.indeterminate_count}",
        f"Cas bloques: {report.executive_summary.blocked_count}",
        (
            "Taux de completude: "
            f"{report.executive_summary.completeness_rate * 100:.2f}%"
        ),
        "",
        "SYNTHESE DES MONTANTS PAR DEVISE ET CERTITUDE",
        *tuple(
            (
                f"{item.currency} | {ras_report_certainty_label(item.certainty.value)} "
                f"| {item.case_count} cas | théorique {item.expected_amount} | "
                f"comptabilisé {item.recorded_amount} | écart {item.difference}"
            )
            for item in report.amount_summaries
        ),
        "",
        "PRINCIPAUX BLOCAGES",
        *tuple(
            f"{ras_report_fact_label(code)} : {count}"
            for code, count in report.executive_summary.top_blockers
        ),
        "",
        "Ce document est une aide a la revue humaine, pas une decision fiscale.",
        "Le détail opérationnel est disponible dans l'export Excel.",
    )
    wrapped_lines = tuple(
        segment
        for line in lines
        for segment in (tuple(wrap(line, width=86)) if line else ("",))
    )
    return _bounded(_pdf_document(wrapped_lines))


def _write_summary(sheet: object, report: RasAuditReport) -> None:
    rows: list[tuple[object, ...]] = [
        ("Rapport d'audit RAS",),
        ("Date de génération", report.generated_at.isoformat()),
        ("Dossiers analysés", report.case_count),
        ("Dossiers conclus", report.executive_summary.concluded_count),
        ("Dossiers probables", report.executive_summary.probable_count),
        ("Dossiers indéterminables", report.executive_summary.indeterminate_count),
        ("Dossiers incomplets", report.executive_summary.blocked_count),
        ("Dossiers en erreur", report.executive_summary.error_count),
        (
            "Anomalies à examiner",
            sum(
                detail.review_priority == "high" or bool(detail.issues)
                for detail in report.details
            ),
        ),
        ("Taux de complétude", report.executive_summary.completeness_rate),
        (),
        (
            "Devise",
            "Certitude",
            "Nombre de cas",
            "RAS theorique",
            "RAS comptabilisee",
            "Ecart comptabilise - theorique",
        ),
    ]
    rows.extend(
        (
            item.currency,
            ras_report_certainty_label(item.certainty.value),
            item.case_count,
            item.expected_amount,
            item.recorded_amount,
            item.difference,
        )
        for item in report.amount_summaries
    )
    rows.extend(
        [
            (),
            ("Devise", "Certitude", "Type d'ecart", "Nombre de cas", "Montant"),
        ]
    )
    rows.extend(
        (
            item.currency,
            item.certainty.value,
            "insuffisance",
            item.insufficient_case_count,
            item.insufficient_amount,
        )
        for item in report.variance_summaries
        if item.insufficient_case_count
    )
    rows.extend(
        (
            item.currency,
            item.certainty.value,
            "excedent",
            item.excess_case_count,
            item.excess_amount,
        )
        for item in report.variance_summaries
        if item.excess_case_count
    )
    for row in rows:
        sheet.append(row)  # type: ignore[attr-defined]
    _style_sheet(sheet)


def _write_detail_sheet(
    workbook: Workbook,
    name: str,
    details: tuple[RasAuditReportDetail, ...],
) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(DETAIL_HEADERS)
    for detail in details:
        sheet.append(_detail_row(detail))
    _style_sheet(sheet, freeze="A2", autofilter=True)


def _detail_row(detail: RasAuditReportDetail) -> tuple[object, ...]:
    return (
        detail.period,
        _safe(detail.document_reference or "Non renseignée"),
        _safe(detail.account_number or ""),
        _safe(ras_report_operation_label(detail.operation_nature)),
        _safe(detail.supplier_reference or "Non renseigné"),
        detail.status_label,
        ras_report_certainty_label(detail.certainty.value),
        ras_report_priority_label(detail.review_priority),
        ras_report_action_label(detail.action_code),
        _safe(detail.explanation),
        detail.tax_base_amount,
        detail.rate_percent,
        detail.expected_amount,
        detail.recorded_amount,
        detail.difference,
        detail.currency,
        _safe("; ".join(detail.missing_fact_labels)),
        _safe("; ".join(detail.issue_labels)),
        _safe("; ".join(detail.legal_source_locators)),
    )


def _style_sheet(
    sheet: object,
    *,
    freeze: str | None = None,
    autofilter: bool = False,
) -> None:
    if freeze is not None:
        sheet.freeze_panes = freeze  # type: ignore[attr-defined]
    if autofilter:
        sheet.auto_filter.ref = sheet.dimensions  # type: ignore[attr-defined]
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in sheet[1]:  # type: ignore[index]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):  # type: ignore[attr-defined]
        width = min(
            45,
            max(
                12,
                *(
                    len(str(cell.value or ""))
                    for cell in sheet[get_column_letter(column)]  # type: ignore[index]
                ),
            ),
        )
        sheet.column_dimensions[get_column_letter(column)].width = width + 2  # type: ignore[attr-defined]


def _pdf_document(lines: tuple[str, ...]) -> bytes:
    pages = tuple(lines[index : index + 48] for index in range(0, len(lines), 48))
    page_ids = tuple(4 + index * 2 for index in range(len(pages)))
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: (
            f"<< /Type /Pages /Kids [{' '.join(f'{item} 0 R' for item in page_ids)}] "
            f"/Count {len(page_ids)} >>"
        ).encode("ascii"),
        3: (
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
            b"/Encoding /WinAnsiEncoding >>"
        ),
    }
    for index, page_lines in enumerate(pages):
        page_id = page_ids[index]
        content_id = page_id + 1
        stream_lines = ["BT", "/F1 10 Tf", "50 790 Td", "14 TL"]
        for line in page_lines:
            stream_lines.extend((f"({_pdf_escape(line)}) Tj", "T*"))
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode("cp1252", errors="replace")
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode("ascii")
        objects[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
            + stream
            + b"\nendstream"
        )
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id in range(1, max(objects) + 1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode("ascii"))
        output.extend(objects[object_id])
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _safe(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _bounded(content: bytes) -> bytes:
    if len(content) > RAS_REPORT_EXPORT_MAX_BYTES:
        raise ValueError("RAS report export exceeds the configured size limit")
    return content
