from dataclasses import replace
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.ras_audit.accounting_assessment import (
    RasAccountingAssessment,
    RasAccountingAssessmentStatus,
)
from app.ras_audit.audit_report import (
    RasAuditReportCase,
    RasAuditReportGenerator,
    RasReportCertainty,
    ras_report_business_payload,
)
from app.ras_audit.report_exports import (
    RAS_REPORT_EXPORT_MAX_BYTES,
    _bounded,
    render_ras_report_pdf,
    render_ras_report_xlsx,
)
from app.ras_audit.rule_resolution import (
    RasRuleEvidence,
    RasRuleResolution,
    RasRuleResolutionStatus,
)
from app.ras_audit.theoretical_calculation import (
    RasTheoreticalCalculation,
    RasTheoreticalCalculationStatus,
)

SOURCE_HASH = "a" * 64


def test_report_separates_currencies_and_certainty_levels() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(
            _case("CANDIDATE-XOF", "XOF", "5000", "4000", complete=True),
            _case("CANDIDATE-USD", "USD", "200", "200", complete=True),
            _case("CANDIDATE-POTENTIAL", "XOF", "1000", "1000", potential=True),
            _case("CANDIDATE-UNKNOWN", None, None, None, complete=False),
        ),
        reference_versions=("rules-v1", "policy-v1"),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert report.case_count == 4
    assert dict(report.certainty_counts) == {
        "indeterminate": 1,
        "potential": 1,
        "supported_provisional": 2,
    }
    summaries = {
        (summary.certainty, summary.currency): summary
        for summary in report.amount_summaries
    }
    assert (
        summaries[
            (RasReportCertainty.SUPPORTED_PROVISIONAL, "XOF")
        ].difference
        == -1000
    )
    assert summaries[(RasReportCertainty.SUPPORTED_PROVISIONAL, "USD")].difference == 0
    assert summaries[(RasReportCertainty.POTENTIAL, "XOF")].case_count == 1
    assert len(summaries) == 3
    assert report.contract_version == "2.0.0"
    assert report.executive_summary.concluded_count == 2
    assert report.executive_summary.probable_count == 1
    assert report.executive_summary.indeterminate_count == 1
    assert report.executive_summary.completeness_rate == Decimal("0.5")
    assert report.details[0].rounding_policy == "none_exact_decimal"
    assert report.details[0].status_label
    assert report.details[0].action_code
    recorded_summaries = {
        (summary.certainty, summary.currency): summary
        for summary in report.recorded_amount_summaries
    }
    assert (
        recorded_summaries[
            (RasReportCertainty.SUPPORTED_PROVISIONAL, "XOF")
        ].recorded_amount
        == 4000
    )
    assert (
        recorded_summaries[
            (RasReportCertainty.POTENTIAL, "XOF")
        ].recorded_amount
        == 1000
    )
    variance = {
        (summary.certainty, summary.currency): summary
        for summary in report.variance_summaries
    }
    xof_variance = variance[(RasReportCertainty.SUPPORTED_PROVISIONAL, "XOF")]
    assert xof_variance.insufficient_case_count == 1
    assert xof_variance.insufficient_amount == 1000
    assert xof_variance.excess_case_count == 0


def test_report_summarizes_recorded_amounts_even_without_expected_amount() -> None:
    detail = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(_case("CANDIDATE-OPEN", "XOF", None, "2500", complete=False),),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert detail.amount_summaries == ()
    assert detail.recorded_amount_summaries[0].recorded_amount == 2500


def test_insufficiencies_and_excesses_do_not_compensate() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(
            _case("INSUFFICIENT", "XOF", "5000", "4000", complete=True),
            _case("EXCESS", "XOF", "5000", "6000", complete=True),
        ),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert report.amount_summaries[0].difference == 0
    variance = report.variance_summaries[0]
    assert variance.insufficient_case_count == 1
    assert variance.insufficient_amount == 1000
    assert variance.excess_case_count == 1
    assert variance.excess_amount == 1000


def test_report_identifier_is_reproducible_and_order_independent() -> None:
    cases = (
        _case("CANDIDATE-B", "XOF", "5000", "5000", complete=True),
        _case("CANDIDATE-A", "XOF", "1000", "0", complete=True),
    )
    generator = RasAuditReportGenerator()

    first = generator.generate(
        source_sha256=SOURCE_HASH,
        cases=cases,
        reference_versions=("policy-v1", "rules-v1"),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    second = generator.generate(
        source_sha256=SOURCE_HASH,
        cases=tuple(reversed(cases)),
        reference_versions=("rules-v1", "policy-v1"),
        generated_at=datetime(2026, 2, 1, tzinfo=UTC),
    )

    assert first.report_id == second.report_id
    assert tuple(detail.candidate_id for detail in first.details) == (
        "CANDIDATE-A",
        "CANDIDATE-B",
    )


def test_report_identifier_changes_with_trace_or_completeness() -> None:
    generator = RasAuditReportGenerator()
    original = generator.generate(
        source_sha256=SOURCE_HASH,
        cases=(_case("CANDIDATE-A", "XOF", "5000", "5000", complete=True),),
        reference_versions=("rules-v1", "policy-v1"),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    changed_detail = replace(
        original.details[0],
        issues=("new_accounting_issue",),
        basis_is_complete=False,
    )
    changed = generator.generate_from_details(
        source_sha256=SOURCE_HASH,
        details=(changed_detail,),
        reference_versions=("rules-v1", "policy-v1"),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert original.report_id != changed.report_id


def test_json_and_csv_exports_are_business_facing_without_internal_codes() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(_case("HASHED-CANDIDATE", "XOF", "5000", "4000", complete=True),),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    generator = RasAuditReportGenerator()

    json_export = generator.to_json(report)
    csv_export = generator.to_csv(report)

    assert "HASHED-CANDIDATE" not in json_export
    assert "article 207" in json_export
    assert "HASHED-CANDIDATE" not in csv_export
    assert "RAS théorique" in csv_export
    assert "Action recommandée" in csv_export
    assert "candidate_id" not in json_export + csv_export
    assert "rule_id" not in json_export + csv_export
    assert "action_code" not in json_export + csv_export
    assert "SYN-TIERS" not in json_export + csv_export
    assert "Honoraires confidentiels" not in json_export + csv_export


def test_authenticated_preview_can_expose_real_document_reference() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(_case("CANDIDATE-1", "XOF", "5000", "4000", complete=True),),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    report = replace(
        report,
        details=(
            replace(report.details[0], document_reference="FAC-2026-00458"),
        ),
    )

    public_payload = ras_report_business_payload(report)
    preview_payload = ras_report_business_payload(
        report,
        expose_document_references=True,
    )

    public_cases = public_payload["cases"]
    preview_cases = preview_payload["cases"]
    assert isinstance(public_cases, list) and isinstance(public_cases[0], dict)
    assert isinstance(preview_cases, list) and isinstance(preview_cases[0], dict)
    assert str(public_cases[0]["document_reference"]).startswith("PIECE-")
    assert preview_cases[0]["document_reference"] == "FAC-2026-00458"


def test_csv_neutralizes_spreadsheet_formula_prefixes() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(_case("=DANGEROUS", "XOF", "5000", "4000", complete=True),),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    report = replace(
        report,
        details=(replace(report.details[0], account_number="=DANGEROUS"),),
    )
    exported = RasAuditReportGenerator().to_csv(report)

    assert "'=DANGEROUS" in exported
    assert "\n=DANGEROUS" not in exported


def test_excel_export_has_operational_sheets_and_safe_cells() -> None:
    generated = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(
            _case("=DANGEROUS", "XOF", "5000", "4000", complete=True),
            _case("OPEN", None, None, None, complete=False),
        ),
        reference_versions=("rules-v1", "policy-v1"),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    generated = replace(
        generated,
        details=(
            replace(generated.details[0], document_reference="FAC-REELLE-2026-001"),
            generated.details[1],
        ),
    )
    workbook = load_workbook(BytesIO(render_ras_report_xlsx(generated)))

    assert workbook.sheetnames == [
        "Synthese",
        "Anomalies",
        "Dossiers incomplets",
    ]
    assert "candidate_id" not in {
        cell.value for sheet in workbook for row in sheet for cell in row
    }
    summary_values = {
        summary[0].value: summary[1].value
        for summary in workbook["Synthese"]
        if len(summary) > 1
    }
    assert summary_values["Dossiers analysés"] == 2
    assert "contract_version" not in summary_values
    exported_values = {
        cell.value for sheet in workbook for row in sheet for cell in row
    }
    assert "FAC-REELLE-2026-001" in exported_values
    assert not any(
        isinstance(value, str) and value.startswith("PIECE-")
        for value in exported_values
    )


def test_pdf_export_is_readable_summary_without_candidate_details() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(_case("PRIVATE-CANDIDATE", "XOF", "5000", "4000", complete=True),),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    reader = PdfReader(BytesIO(render_ras_report_pdf(report)))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert "RAPPORT DE SYNTHESE - AUDIT RAS" in text
    assert "Dossiers analysés : 1" in text
    assert "PRIVATE-CANDIDATE" not in text


def test_empty_report_has_zero_completeness_and_valid_exports() -> None:
    report = RasAuditReportGenerator().generate(
        source_sha256=SOURCE_HASH,
        cases=(),
        reference_versions=("rules-v1",),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert report.executive_summary.completeness_rate == 0
    assert render_ras_report_xlsx(report)
    assert render_ras_report_pdf(report)


def test_report_export_rejects_content_above_explicit_size_limit() -> None:
    with pytest.raises(ValueError, match="size limit"):
        _bounded(b"x" * (RAS_REPORT_EXPORT_MAX_BYTES + 1))


def test_report_rejects_duplicate_candidates() -> None:
    case = _case("DUPLICATE", "XOF", "5000", "5000", complete=True)

    with pytest.raises(ValueError, match="duplicate"):
        RasAuditReportGenerator().generate(
            source_sha256=SOURCE_HASH,
            cases=(case, case),
            reference_versions=("rules-v1",),
        )


def _case(
    candidate_id: str,
    currency: str | None,
    expected: str | None,
    recorded: str | None,
    *,
    complete: bool = False,
    potential: bool = False,
) -> RasAuditReportCase:
    expected_amount = Decimal(expected) if expected is not None else None
    recorded_amount = Decimal(recorded) if recorded is not None else None
    difference = (
        recorded_amount - expected_amount
        if expected_amount is not None and recorded_amount is not None
        else None
    )
    status = (
        RasAccountingAssessmentStatus.POTENTIAL_RELATED_ENTRY
        if potential
        else (
            RasAccountingAssessmentStatus.PROVISIONAL_RECONCILED
            if difference == 0
            else (
                RasAccountingAssessmentStatus.PROVISIONAL_AMOUNT_MISMATCH
                if difference is not None
                else RasAccountingAssessmentStatus.INDETERMINATE
            )
        )
    )
    assessment = RasAccountingAssessment(
        status=status,
        candidate_entry_id=candidate_id,
        expected_amount=expected_amount,
        recorded_amount=recorded_amount,
        difference=difference,
        currency=currency,
        tolerance=Decimal("0") if complete else None,
        policy_version="policy-v1" if complete else None,
        missing_facts=() if complete else ("confirmed_document_link",),
        issues=(),
        potential_adjustments_present=False,
        basis_is_complete=complete,
    )
    resolution = RasRuleResolution(
        status=RasRuleResolutionStatus.RESOLVED_PROVISIONAL,
        rule_id="RULE-1",
        rule_version="rules-v1",
        calculation_method="flat_rate",
        rate_percent=Decimal("5"),
        missing_facts=(),
        alternative_rule_ids=(),
        evidence=(
            RasRuleEvidence(
                source_kind="rate",
                source_url="https://example.invalid/legal-source",
                source_locator="article 207",
                source_sha256="b" * 64,
            ),
        ),
        source_assurance="synthetic",
        fact_sources=(("currency", "gl"),),
    )
    calculation = RasTheoreticalCalculation(
        status=(
            RasTheoreticalCalculationStatus.CALCULATED_PROVISIONAL
            if expected_amount is not None
            else RasTheoreticalCalculationStatus.NOT_CALCULABLE
        ),
        rule_id="RULE-1" if expected_amount is not None else None,
        rule_version="rules-v1" if expected_amount is not None else None,
        base_fact_name="tax_base_amount" if expected_amount is not None else None,
        base_amount=expected_amount,
        expected_amount=expected_amount,
        currency=currency,
        calculation_method="flat_rate" if expected_amount is not None else None,
        rate_percent=Decimal("5") if expected_amount is not None else None,
        steps=(),
        parameter_versions=(),
        source_locators=("article 207",) if expected_amount is not None else (),
        rounding_policy="none_exact_decimal",
        reason=None if expected_amount is not None else "missing_facts",
    )
    return RasAuditReportCase(
        candidate_id=candidate_id,
        assessment=assessment,
        resolution=resolution,
        calculation=calculation,
    )
